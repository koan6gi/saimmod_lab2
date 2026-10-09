import os
import sys
import random
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from src.config import Config, load_config
from src.simulation import SupermarketSimulation


def create_excel_report(output_file: str = "supermarket_simulation.xlsx", num_replications: int = 1000):
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    sub_header_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    sub_header_font = Font(name="Calibri", size=10, bold=True, color="1E3A8A")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # ==========================================
    # SHEET 1: input (Входные данные модели)
    # ==========================================
    ws_input = wb.create_sheet(title="input")
    ws_input.views.sheetView[0].showGridLines = True

    # Headers for 3 sections as in process.xlsx
    ws_input["A1"] = "Run input data"
    ws_input["E1"] = "Controlled input data (parameters)"
    ws_input["H1"] = "Uncontrolled input data (variables)"

    for cell_ref in ["A1", "E1", "H1"]:
        ws_input[cell_ref].font = header_font
        ws_input[cell_ref].fill = header_fill
        ws_input[cell_ref].alignment = Alignment(horizontal="center", vertical="center")

    ws_input.merge_cells("A1:C1")
    ws_input.merge_cells("E1:G1")
    ws_input.merge_cells("H1:J1")

    # Rows content
    input_rows = [
        # (A, B, C, E, F, G, H, I, J)
        ("Number of replications", num_replications, "", "Number of departments", 5, "pcs", "Arrival interval (mean)", 1.0, "min"),
        ("Time of run", 60.0, "min", "Sellers per department", 1, "pcs", "Arrival distribution", "exponential", ""),
        ("with clearance", "yes", "", "Number of buyers at start", 2, "buyers", "Service & payment time (mean)", 1.8, "min"),
        ("is logged", "no", "", "Min time of payment", 0.4, "min", "Service & payment std", 0.5, "min"),
        ("RNG seed", 42, "", "Shopping departments basket", "1..5", "depts", "Service distribution", "normal (truncated)", ""),
        ("", "", "", "Discipline of queues", "FIFO", "", "Triangular mode of basket", 3, "depts")
    ]

    for row_idx, data in enumerate(input_rows, start=2):
        # Run input data
        ws_input[f"A{row_idx}"] = data[0]
        ws_input[f"B{row_idx}"] = data[1]
        ws_input[f"C{row_idx}"] = data[2]
        # Controlled input data
        ws_input[f"E{row_idx}"] = data[3]
        ws_input[f"F{row_idx}"] = data[4]
        ws_input[f"G{row_idx}"] = data[5]
        # Uncontrolled input data
        ws_input[f"H{row_idx}"] = data[6]
        ws_input[f"I{row_idx}"] = data[7]
        ws_input[f"J{row_idx}"] = data[8]

        for col_letter in ["A", "B", "C", "E", "F", "G", "H", "I", "J"]:
            c = ws_input[f"{col_letter}{row_idx}"]
            c.font = data_font
            c.border = thin_border
            if col_letter in ["A", "E", "H"]:
                c.fill = sub_header_fill
                c.font = sub_header_font

    # ==========================================
    # SHEET 2: responses (Результаты 1000 прогонов)
    # ==========================================
    ws_resp = wb.create_sheet(title="responses")
    ws_resp.views.sheetView[0].showGridLines = True

    resp_headers = [
        "run number", "time of run", "number of buyers", 
        "avg time in supermarket", "avg total queue length", "avg busy sellers"
    ]
    ws_resp.append(resp_headers)
    for col_idx in range(1, len(resp_headers) + 1):
        cell = ws_resp.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    print(f"Запуск {num_replications} репликаций для таблицы Excel...")
    cfg_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "config.json"))
    base_config = load_config(cfg_path)
    base_config.simulation.trace = False
    
    first_sim_discrete_history = []
    first_sim_continuous_history = []

    for r in range(1, num_replications + 1):
        cfg = load_config(cfg_path)
        cfg.simulation.seed = 42 + r * 101
        cfg.simulation.trace = False
        sim = SupermarketSimulation(cfg)
        res = sim.run()

        ws_resp.append([
            r,
            round(res["simulation_end_time"], 3),
            res["total_served"],
            round(res["avg_time_in_supermarket"], 3),
            round(res["avg_total_queue_len"], 3),
            round(res["avg_busy_sellers"], 3)
        ])

        if r == 1:
            # Сохраняем подробную историю первого прогона для листов динамики
            first_sim_discrete_history = list(zip(
                sim.stats.time_in_supermarket.timestamps,
                sim.stats.time_in_supermarket.observations,
                sim.stats.time_in_supermarket.running_means
            ))
            first_sim_continuous_history = sim.stats.total_queue_len.history

    # ==========================================
    # SHEET 3: time in store_dyn (Дискретный отклик)
    # ==========================================
    ws_time = wb.create_sheet(title="time in store_dyn")
    ws_time.views.sheetView[0].showGridLines = True
    ws_time.append(["replication", "modelTime", "value", "movingAvg"])
    for col_idx in range(1, 5):
        cell = ws_time.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for t_stamp, val, r_mean in first_sim_discrete_history:
        ws_time.append([1, round(t_stamp, 3), round(val, 3), round(r_mean, 3)])

    # ==========================================
    # SHEET 4: queue length_dyn (Непрерывный отклик)
    # ==========================================
    ws_queue = wb.create_sheet(title="queue length_dyn")
    ws_queue.views.sheetView[0].showGridLines = True
    ws_queue.append(["replication", "modelTime", "value", "movingAvg"])
    for col_idx in range(1, 5):
        cell = ws_queue.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for tp in first_sim_continuous_history:
        ws_queue.append([1, round(tp.time, 3), round(tp.instant_value, 3), round(tp.cumulative_average, 3)])

    # Auto-adjust column widths for all sheets
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(output_file)
    print(f"Excel-файл успешно сформирован и сохранен в: {output_file}")


if __name__ == "__main__":
    default_out = os.path.abspath(os.path.join(os.path.dirname(__file__), "supermarket_simulation.xlsx"))
    create_excel_report(default_out, num_replications=1000)
